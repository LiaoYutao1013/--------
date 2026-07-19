from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('04_Papers_&_Literature (论文精读)') / '无人方程式赛车论文感知子系统综述.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)

# source | title | overview | perception | takeaway
DATA = r'''
00简介.pdf|前置知识与资料索引（参考材料）|面向车队论文阅读的中文前置笔记，覆盖组合惯导、里程计、线控、SLAM、雷达、标定及若干规划控制摘要。|明确 INS 参考点、相机内外参、点云投影、视觉里程计、后端优化和回环检测的边界，是后续材料的术语字典。|工程重点是统一坐标系、时间戳和安装后的标定状态；不能把概念笔记替代为可运行的感知方案。
01苏黎世车队无人车论文2019.pdf|AMZ Driverless: The Full Autonomous Racing System|AMZ/ETH 给出覆盖感知、估计、SLAM、规划、控制与测试的完整赛车软件栈，并在未知赛道竞赛验证。|独立 LiDAR、视觉及双模态管线输出锥桶三维位置和颜色；LiDAR 去地面聚类，三台全局快门相机以双目近距和单目远距互补；五类传感器融合后进入粒子滤波 SLAM。|冗余应能独立退化运行，并将量测置信度、失效检测和低延迟当作规划性能的一部分。
02北理工论文2017.pdf|Autonomous Driving System Design for Formula Student Driverless Racecar|北理工 Smart Shark I 的整车论文，介绍 ROS 上位机、低层 VCU、检测、建图和轨迹跟踪闭环。|16 线 LiDAR 提取锥桶候选和几何位置，相机分类颜色；GPS-INS 与 LiDAR 里程计耦合成高频定位，并累积彩色锥桶地图。|ROS 数据流要显式处理时间同步和坐标变换，实时执行与安全功能宜保留在 VCU。
03匈牙利 BME Formula Racing Team车队论文2025.pdf|The Autonomous Software Stack of the FRED-003C|BME 总结从学生项目到全尺寸无人赛车的栈，按状态估计、感知、规划、控制划分高算力和实时层。|车头 LiDAR 提供独立锥桶空间观测，广角相机用于颜色分类；通过标定对齐图像和点云，并使用 LiDAR-惯性里程计强化状态估计。|每项观测应带时间、坐标、协方差和健康状态，使后端能拒绝坏量测。
04印度车队（无价值）.pdf|IIT Bombay Racing Driverless: Autonomous Driving Stack for Formula Student AI|IIT Bombay 的原型栈使用 Jetson Orin、两台 ZED2i、VLP-16 和 GNSS/INS，在小车及仿真验证。|LiDAR 和双目并行：LiDAR 提供较准空间候选，双目覆盖远距检测和深度；采用 RANSAC 去地面，并进行图像-点云对应。|融合链路有原型价值，但要在赛车前量化远距深度、标定、漏检和端到端延迟。
05奥地利维亚纳理工（一般）.pdf|Design of an Autonomous Race Car for Formula Student Driverless|TU Wien Racing 的早期整车设计，讨论从人工赛车改装到无人系统的框架、传感器、检测和建图。|ZED 双目与激光扫描器并行检测锥桶；融合考虑各传感器可靠性，SLAM 依托锥桶观测和运动信息，在短暂失明下维持地图。|不要假设任一传感器常可用；外参、重复观测、颜色不确定性都需进入地图更新。
06 德国KIT 2022  .pdf|The Software Stack That Won the Formula Student Driverless Competition|KA-RaceIng 获奖栈在未知赛道实现高速闭环，采用高精度地图、赛线优化与 MPC。|360 度 LiDAR 为主、三相机可选，约 35 m 稳定识别锥桶；GraphSLAM 建图，报告高速窄赛道地图 RMSE 小于 15 cm，三角剖分和样条恢复边界。|感知范围、地图误差和车速必须共同设计；略增可见距离也会改变可规划视野。
07 AMZ 2020.pdf|Accurate Mapping and Planning for Autonomous Racing|AMZ 聚焦未知赛道的感知-建图-规划联动，目标是在首圈更快行驶且保持地图质量。|相机与 LiDAR 早期融合检测锥桶，卡尔曼滤波时间融合；分层地图允许规划先使用及时低置信地图，再沉淀高质量地图。论文报告未知环境可靠速度由 3 m/s 提至 12 m/s、地图 RMSE 约 0.29 m。|区分快速可用地图和精确持久地图，保存观测来源、置信度和更新历史。
08 北理工 2018 .pdf|From Perception to Control: An Autonomous Driving System for a Formula Student Driverless Car|Smart Shark II 展示冗余感知、EKF、占据栅格与 MPC 的整车链路。|3D LiDAR、单目、GPS/INS 与轮速协同；LiDAR 提供空间候选，视觉完成颜色与标定位姿，GNSS/INS 和 LiDAR 里程计经 EKF 融合，结果写入占据栅格。|应同时建设检测冗余与状态估计冗余，并把单模态可用性、融合置信度发给规划层。
09 德国KIT 2020.pdf|The Autonomous Racing Software Stack of the KIT19d|KIT19d 强调模块化、可靠性、快速开发和性能之间的平衡。|四 LiDAR、三相机形成前后覆盖；同步 ECU 触发并拼接点云，点云产生地标候选，图像验证候选并补颜色，IMU/轮速/地标进入定位建图。|硬件触发、时间关系、外参版本和重叠区一致性应列为感知验收项。
10 澳洲QUT 2023.pdf|Racing With ROS 2: A Navigation System for an Autonomous Formula Student Race Car|QUT 比较自研栈和 ROS 2 现成导航组件，重在可复用的系统集成。|双目、LiDAR、INS 输入经深度锥桶检测、点云聚类、EKF、slam_toolbox 等组件组合，输出地图和路径接口。|先稳定消息语义、坐标、频率和延迟，再替换算法；开源组件仍需在高速条件复测。
11 德国奥格斯堡 2024.pdf|Lane Detection using Graph Search and Geometric Constraints for Formula Student Driverless|论文从含大量误检的二维锥桶点集中恢复长距离车道，提出带学习似然和几何约束的 CLC 回溯图搜索。|把前端误检和有限视野作为输入现实；报告低于 15 ms、可检测 100 m 以上车道，面对 50% 假阳性仍具鲁棒性。|可借助赛道几何先验消歧，不必等待零误检前端；需联合使用前端置信度和速度相关可见距离。
12 葡萄牙里斯本高等高等理工学院 2021.pdf|Path Planning and Guidance Laws of a Formula Student Driverless Car|短文主攻人工势场路径、两遍速度规划和解耦控制，验证主要在仿真。|感知仅作为上游假设：GPS、视觉及融合输出赛道边界和静态障碍，规划据此生成路径。|其价值在于反推感知需求：边界刷新率、障碍位置和置信度必须支持保守速度策略。
13葡萄牙里斯本高等高等理工学院 2021 .pdf|A SLAM Method for the Formula Student Driverless Competition|论文比较两种粒子滤波 SLAM 和图优化 SLAM，并提出结合跟踪信息的数据关联方案。|LiDAR、RGB 相机及姿态速度信息提供带颜色锥桶观测；重点处理关联、延迟补偿、地图更新和建图后切换定位。|检测结果不等于地图；后端需要时间、类别、位置协方差及轨迹上下文。
14 慕尼黑TUM 2022.pdf|Autonomous Racing: A Survey on Perception, Planning and Control|TUM 综述自动赛车感知、规划、控制和端到端学习，是检索入口而非单一方案。|感知涵盖目标/自由空间、定位、建图和状态估计，归纳高速低延迟、精确定位、宽域定位和极端工况挑战，涉及相机、LiDAR、GNSS、IMU、轮速融合。|评价不要只看单帧精度，应以范围、延迟、失效概率及圈速/安全影响闭环衡量。
15 都灵理工大学 2021.pdf|Global Mapping Algorithm for a Driverless Race Car|都灵理工硕士论文围绕从局部观测到全局锥桶地图的 ROS 算法。|ZED 双目输出锥桶位置与颜色，Velodyne VLP-16 提供独立位置，SBG Ellipse-N 提供惯导定位；测试主要用双目，LiDAR 作为增强来源，强调时间戳同步。|采用单传感器可用、融合后增强的地图设计，并补充实时融合和地图质量量化。
16 德国拜罗伊特大学2024 .pdf|Winning Through Simplicity: Autonomous Car Design for Formula Student|Elefant Racing 面向小团队，强调简单可维护的整车设计。|近距宽视场和远距窄视场单目相机，探索双目与 Flash LiDAR；旋转 LiDAR 扫描小锥桶导致早融合困难，因此讨论独立检测后的晚期融合；使用同步触发和棋盘格标定。|高速小目标下时间对齐不足时，晚期融合可能比强行早融合更稳。
17 葡萄牙里斯本高等高等理工学院 2022.pdf|Path Planning and Guidance Laws of a Formula Student Driverless Car（期刊版）|同题研究的期刊扩展版，完整给出人工势场、速度轮廓、障碍规避和解耦控制。|感知仍为输入前提，未提出锥桶检测、标定或 SLAM 新方法。|可形成感知需求表：边界精度与刷新率影响势场，障碍物置信度影响安全距离与速度上限。
18 葡萄牙里斯本高等高等理工学院 2021.pdf|Path Planning and Guidance Laws of a Formula Student Driverless Car（学位论文版）|该长篇学位论文扩展了同一主题的路径、控制和仿真实现。|将 LiDAR、相机、GPS 与融合定位视为赛道地标和位姿的上游输入，未给出前端模型。|仿真应注入漏检、延迟、偏差与颜色错分，检验规划而非默认完美感知。
19 西班牙Universitat Politècnica de Catalunya 2020 .pdf|Formula Student Driverless: The autonomous systems in the Skidpad event|论文针对八字 Skidpad 分析 SLAM、路径、控制和任务流程。|使用一 LiDAR 和两相机观察锥桶，讨论只看到单侧、单个或多组锥桶等不完整观测，并用任务状态判断圈次位置。|Skidpad 可凭强几何先验补偿缺失；Autocross 则更依赖持续建图和大前视距离。
20 维也纳技术大学 2017.pdf|Path Planning and Control in an Autonomous Formula Student Vehicle|Monash Motorsport 研究未知锥桶赛道建图、最快路径、MPC 与赛线优化。|感知输入是持续更新的锥桶位置，系统在视野内发现锥桶并累积赛道地图；附件未给出传感器型号和检测细节。|接口应提供持续、去重、坐标一致的地标流与可见范围，避免暂时漏检被误解为赛道终点。
2024FSAC入门指南.pdf|大学生方程式无人驾驶系统入门指南 2024（参考材料）|规则安全与车检指南，涵盖 RES、SDC、ASMS、SCS、ASB、状态机、安装和数据记录。|感知信号丢失、异常或上层算法不能保证安全时，应作为系统关键信号触发安全状态和 EBS；安装需考虑干扰和机械风险。|感知是安全关键输入，每条链路需在线健康监测、超时/范围检查、日志和降级策略。
21 意大利帕多瓦大学 2023 .pdf|Design of a perception system for the Formula Student Driverless competition: from vehicle sensorization to SLAM|帕多瓦硕士论文从传感器化、锥桶感知到因子图 SLAM，服务 RaceUP 原型建设。|双目与 LiDAR 两条冗余检测支路，视觉采用 YOLO 类检测；后端以车辆位姿和地标构造因子图，利用空间和运动学约束优化。|先完成视场和安装设计，再独立验证前端，最后以带协方差观测接入因子图。
22 挪威科技大学 2024 .pdf|Application of Koopman Operator Theory for Control of an Autonomous Formula Student Race Car|NTNU 研究 Koopman 提升模型和预测控制，核心结论来自仿真。|系统介绍中鼻端 Ouster OS1-32G 提取锥桶中心，图优化 SLAM 融合锥桶与运动信息，建图后可切换仅定位；感知不是创新核心。|控制模型必须获得稳定、可解释的状态；要分辨感知噪声与控制模型误差。
MPCC局部规划与控制_无人驾驶赛车行车路径规划研究_龚国铮.pdf|无人驾驶赛车行车路径规划研究|本科论文研究 Delaunay、贪婪中心线、赛道宽度优化及 MPCC 规划控制。|感知只以桩桶点集作为输入；颜色缺失时几何代价仍有一定容错，属于规划对感知误差的鲁棒性。|接口应含位置置信度和可能类别；输入质量下降时规划应收缩可行域并降速。
Toward_Fair_and_Thrilling_Autonomous_Racing_Governance_Rules_and_Performance_Metrics_for_the_Autonomous_One (1).pdf|Toward Fair and Thrilling Autonomous Racing: Governance Rules and Performance Metrics|赛事治理与性能指标短文，讨论 Roborace、IAC、FSD 的公平性和教育目标。|不提出车载感知算法，但将感知、定位、预测、规划控制并列为栈能力，障碍避让和超车直接提升检测追踪需求。|可转化为检测距离、漏检误检、追踪连续性、端到端延迟和失效安全等验收指标。
TUM autonomous motorsport_ An autonomous racing software for the Indy Autonomous Challenge.pdf|TUM autonomous motorsport: An autonomous racing software for the Indy Autonomous Challenge|附件文本层主要为下载页眉，无法可靠提取正文方法、实验和数值。|只能确认其主题为 IAC 整车软件，不能据此证实传感器、检测、融合或定位方案。|应更换可搜索 PDF 或做高质量 OCR 后重读；本综述不把它作为感知证据来源。
Vol4_04.pdf|er.autopilot 1.0: The Full Autonomous Stack for Oval Racing at High Speeds|TII EuroRacing 面向 IAC 椭圆赛道，覆盖障碍规避、主动超车、超高速控制、仿真、遥测与安全。|双 GNSS、三 LiDAR、六相机、三雷达；RTK GNSS 多天线定位，EKF 融合 GNSS/LiDAR/IMU/轮速，视觉 CNN 检测目标，动态对手需追踪融合。|高速关键是可用距离和端到端时延，应将覆盖、追踪、定位冗余和安全监控统一预算。
无人驾驶电动赛车路径规划与跟踪控制研究.pdf|无人驾驶电动赛车路径规划与跟踪控制研究|辽宁工业大学硕士论文研究定位建图、多赛道规划、跟踪控制和 FSAC 实车验证。|组合惯导给车辆位置，LiDAR 点云得到锥桶相对位置，经坐标变换建全局地图，并实时接入相机、LiDAR、惯导。|统一车辆参考点、时间同步和观测不确定度，地图质量直接决定多赛项规划。
面向无人方程式赛车的道路感知和轨迹规划算法研究_高群森.pdf|面向无人方程式赛车的道路感知和轨迹规划算法研究|长安大学硕士论文直接研究锥桶检测、双目建图、高速轨迹和 ROS 实车验证。|YOLOv7 检测并排序锥桶；双目经 MATLAB 标定，用 SGM 获得深度，再融合组合惯导位姿，建立赛道地图。|按检测、深度、定位、地图逐级验证，并评估地图误差对 Delaunay 中心线及 Frenet/Lattice 轨迹的影响。
'''

# Each entry is deliberately a technical contribution assessment, not an assertion
# that every attached item contains a novel perception algorithm.
INNOVATIONS = r'''
+这份材料本身没有提出可验证的新算法，创新性在于把比赛论文中的传感器、坐标系和系统术语压缩成可复用的学习入口。对工程实践而言，它把“外参-内参-参考点-坐标变换”串成同一条链，能减少后续系统集成中常见的概念错位；应视为知识组织贡献，而非论文级感知创新。
创新点首先是可切换、可独立运行的 LiDAR 与视觉感知链路，再把多源观测接入统一的状态估计和 FastSLAM 地图。它不是只给出某个检测器，而是证明感知、估计、规划和极限控制能在真实未知赛道闭环协同；这种端到端竞赛验证使模块接口、容错和测试框架成为重要贡献。
相较于仅用单一锥桶检测的早期方案，本文将 LiDAR 的几何定位、相机颜色分类、GPS-INS 和 LiDAR 里程计组织为同步 ROS 管线，并在上位机与 VCU 之间分离计算与安全控制。其创新性主要为面向 FSAC 的整车系统集成和可运行验证，而非提出新的通用检测网络。
BME 的贡献是将感知、状态估计、规划、控制与实时执行层解耦，并以独立 LiDAR 观测保证几何主链、相机承担语义补充。论文还展示了从粒子滤波建图向 LiDAR-惯性里程计演化的取舍，强调软件架构和故障隔离如何支撑更大尺度赛车，属于成熟技术的工程化组合创新。
IIT 的技术路线强调三路并行观测：点云支路提供近中距几何，双目支路提供远距视觉与深度，再在相机框和点云之间做关联。用 RANSAC 地面分割降低点云搜索空间，并为未匹配目标保留纯视觉深度后备，这种“互补而非互相依赖”的设计是其主要创新；不过证据主要来自原型与仿真。
该早期方案的亮点是以传感器可靠性为融合依据，而非静态地相信某一路数据；映射模块显式面对相机或激光短时失明。对当时的 FSD 场景，这种将可用性管理写入 SLAM 的系统设计比某个单一视觉模型更有价值，但论文未给出新的概率融合理论或大规模性能比较。
KA-RaceIng 将 360 度 LiDAR 的大视场、可选多相机语义、GraphSLAM 和高性能赛线串成一套能在高速运行的系统。创新的系统指标很明确：约 35 m 锥桶观测支持超过 70 km/h 建图，地图误差控制在 15 cm 量级；其价值在于证明“完整覆盖+图优化地标地图”可以直接换取更大的规划视野和可用速度。
AMZ 的核心技术创新是早期相机-LiDAR 融合与两层地图的组合。前者将图像语义和 LiDAR 深度在检测前结合，后者允许低置信地图立即用于首圈规划、同时用更谨慎的估计沉淀持久地图；这直接解决未知赛道中速度和地图质量互相牵制的问题，并以速度和 RMSE 的对比给出实证。
Smart Shark II 的创新性是把两条独立锥桶检测链、EKF 状态融合和占据栅格累积串成冗余闭环。LiDAR 负责强几何感知，单目相机解决颜色语义，GNSS/INS 与 LiDAR 里程计将低频绝对定位和高频相对运动互补；贡献在于容错整车架构与真实赛事验证，而非全新的单模块算法。
KIT19d 把多 LiDAR 的同步触发、点云拼接和“LiDAR 生成候选、相机验证候选”的级联检测做成完整工程链。相比简单后融合，硬件同步 ECU 与重叠视场处理将时序误差前置解决，使多传感器部署真正能带来覆盖与冗余；论文还提供运行时和竞赛数据，强化了系统可复现性。
QUT 的创新不是重造感知模型，而是评估并组装 ROS 2 生态中已有的 robot_localization、slam_toolbox 和导航组件，用开源成熟模块替代脆弱的自研链路。它的技术贡献是接口层和评测层：说明哪些现成组件可承接双目、LiDAR、INS 数据，以及高速赛车条件下仍需补哪些适配与验证。
CLC 的创新在于把“锥桶点有误检、颜色可能缺失、视野不完整”直接建模为组合搜索问题，而不是把前端输出当作确定边界。它利用学习的似然函数排序候选、用赛道几何约束剪枝、通过回溯图搜索寻找长视距全局解；在低延迟和高假阳性条件下仍有量化结果，是附件中最明确的感知后端算法创新之一。
本文的技术主张在规划控制而非感知：用人工势场构造横向参考、两遍算法生成速度轮廓，并通过横纵解耦控制实现赛道与障碍规避。其与感知的创新关系在于把前端地图质量转换为势场与速度规划的边界条件，因而更适合作为感知需求分析材料，不应归入锥桶检测或融合文献。
相对于仅比较 FastSLAM 变体，论文引入图优化 SLAM 并提出融合跟踪信息的数据关联策略，针对锥桶重复、噪声和延迟观测减少错误匹配。结果显示图优化在精度和计算效率上优于所比较的粒子滤波方案；对感知系统而言，这说明创新重点在“如何消费观测并关联地标”，不在传感器本身。
综述的创新是建立跨赛事、跨尺度的感知问题框架：将物体/自由空间检测、地图、定位和状态估计统一到高速自动赛车约束下比较。它把速度、延迟、定位精度、感知距离和计算资源视作耦合因素，为后续选型和评测提供结构化问题清单，而不宣称提出新的前端算法。
论文的贡献是把视觉锥桶观测、惯导定位和逐步地图更新落实为全局映射算法，并在 ROS 通信与时间同步约束下组织实现。其创新属于“由局部反应式地图向可维护全局地图过渡”的工程实现；将双目作为当前主链、LiDAR 作为独立增强链，也为后续融合迭代留出清晰接口。
该文对小团队最有价值的创新是选择性简化：先用同步相机系统提供可靠视觉支路，而不在早期强行实现难以校准的旋转 LiDAR 早融合。它分析 Flash LiDAR 与全局快门相机在时序上的匹配优势，并提出在传统 LiDAR 条件下采用独立检测和晚期融合，是由传感器物理采样方式驱动的架构决策。
期刊版的主要创新仍是规划控制方法：人工势场负责路径结构，双遍速度规划将曲率和牵引约束转成可行速度，解耦控制负责执行。相较短文，期刊版更系统地将未知静态障碍纳入路径问题；对感知的间接贡献是明确了地图/障碍状态应如何影响速度与安全裕度。
学位论文版将路径规划、指导律、车辆模型和仿真验证展开得更完整，创新主要是针对 Formula Student 平台的方案比较与实现细节。感知部分没有提出新的观测模型，因此其真正用途是为感知团队构建带噪声、漏检和延迟的闭环仿真测试基准，而不是作为前端设计依据。
在结构高度规则的 Skidpad 中，论文将锥桶观测、SLAM 地图与项目状态机结合，利用赛道拓扑判断当前任务阶段。这种把任务先验用于弥补稀疏感知的设计是其特色；它不是通用未知道路感知算法，但对专用赛项能显著降低对全局完整地图的依赖。
论文将实时新锥桶观测持续并入赛道地图，再以车辆模型、自动微分和非线性优化生成最快赛线。创新主要在从增量地标地图到可优化赛线的接口设计，以及面向特定车辆的参数化优化；原文不提供前端细节，所以不能将其优化性能归因于某种传感器算法。
指南的创新性不属于研究算法，而在于把安全回路、遥控急停、状态定义、系统关键信号和车检可测试性整合为可执行工程约束。它把感知失效明确提升为触发紧急制动的功能安全事件，为感知健康监测提供了比“检测精度”更严格的系统定义。
该论文的技术贡献是把传感器选型、双目/点云锥桶观测、数据关联和因子图后端作为一条连续设计链。图优化使用车辆位姿和锥桶地标的空间、运动学约束，而非只依赖滤波递推；这种后端选择有利于在回看和重线性化时消化前端观测的不一致性。
论文的真正创新是控制建模：使用 Koopman 提升、动态模态分解及神经网络提升，将非线性车辆动力学近似成可供线性 MPC 使用的高维线性模型。感知部分只是向控制提供锥桶和状态，价值在于提醒系统必须量化观测噪声，否则模型学习到的误差会被误认为车辆动力学。
本文将 Delaunay 离散、贪婪边界/中心线提取与快速 MPCC 结合：先由粗几何路径产生可行解，再在线优化使轨迹更平滑且更接近可用赛道宽度。论文还进行线性化和失效处理以提升求解速度；其创新在规划控制一体化，对颜色缺失的容错来自几何而非感知学习。
该文提出规则与性能度量层面的创新：通过赛事任务、标准平台、准入规则和性能指标平衡公平性、观赏性与教育价值。它使感知能力可以被赛项目标度量，例如障碍绕行和超车把检测、跟踪、预测从隐含要求变为可比较指标，但并不提供车载算法。
由于附件正文文本层不可用，不能可靠识别论文作者声称的技术创新、实验设置或量化结果。唯一可负责任的创新性判断是：它应是一篇面向 IAC 的整车软件论文，但在取得可读原文前，任何对其感知或定位方法的具体描述都应视为未经证实。
er.autopilot 的创新在于将极高速椭圆赛道的定位、静态障碍规避、动态对手感知、超车规划和安全遥测整合到一辆 IAC 赛车。多 GNSS、多 LiDAR、相机、雷达的异构感知与 EKF/LiDAR 定位共同服务超过 75 m/s 的运行，突出“动态目标跟踪与安全冗余”相对锥桶赛事的新增难题。
论文的技术贡献是将组合惯导车辆状态、LiDAR 锥桶点云、相机数据和不同比赛项目的规划方法连成整车闭环。其映射方法通过位姿与锥桶相对坐标变换构造全局地图，创新更偏向面向 FSAC 的系统实现与实车验证；未提出可与深度检测网络直接比较的新感知算法。
论文的感知创新最明确：以 YOLOv7 构建并训练锥桶数据集，完成检测、定位和排序；再以标定双目和 SGM 立体匹配恢复深度，并与组合惯导位姿融合为世界坐标地图。它把学习检测、几何深度和惯导定位逐层连接到 Delaunay 与 Frenet/Lattice 规划，是一条可分别评估且可实车验证的完整链路。
'''

def font(run, size=11, bold=False, color=(25,25,25), italic=False):
    run.font.name = 'Microsoft YaHei'
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for k in ('ascii', 'hAnsi', 'eastAsia'):
        rf.set(qn('w:' + k), 'Microsoft YaHei')
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run.font.color.rgb = RGBColor(*color)

def field(paragraph, instruction):
    e = OxmlElement('w:fldSimple'); e.set(qn('w:instr'), instruction); paragraph._p.append(e)

def labeled(doc, label, text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.25; p.paragraph_format.keep_together = True
    font(p.add_run(label), 11, True, (31,77,120)); font(p.add_run(text))

items = [tuple(row.split('|')) for row in DATA.strip().splitlines()]
innovations = INNOVATIONS.strip().splitlines()
assert len(items) == len(innovations), '每篇材料必须对应一条创新性判断'
doc = Document(); sec = doc.sections[0]
for a in ('top_margin','bottom_margin','left_margin','right_margin'): setattr(sec, a, Inches(1))
sec.header_distance = Inches(.492); sec.footer_distance = Inches(.492)
normal = doc.styles['Normal']; normal.font.name = 'Microsoft YaHei'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); normal.font.size = Pt(11); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for n, s, c, b, a in [('Heading 1',16,(46,116,181),18,10),('Heading 2',13,(46,116,181),14,7)]:
    st=doc.styles[n]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(s); st.font.color.rgb=RGBColor(*c); st.font.bold=True; st.paragraph_format.space_before=Pt(b); st.paragraph_format.space_after=Pt(a); st.paragraph_format.keep_with_next=True
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.LEFT; font(hp.add_run('无人方程式赛车论文感知子系统综述'),9,True,(90,90,90))
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(fp.add_run('第 '),9,False,(90,90,90)); field(fp,'PAGE'); font(fp.add_run(' 页'),9,False,(90,90,90))
for _ in range(5): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18); font(p.add_run('论文精读 | 感知子系统专题'),12,True,(122,90,0))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8); font(p.add_run('无人方程式赛车论文感知子系统综述'),28,True,(11,37,69))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(30); font(p.add_run('逐篇总结附件论文与参考材料，聚焦传感器、检测融合、定位建图及工程约束'),12,False,(80,80,80))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run('阅读范围：30 份附件 PDF | 生成日期：2026-07-17'),10,False,(80,80,80))
doc.add_page_break(); doc.add_heading('阅读说明',1)
labeled(doc,'目标：','按附件顺序逐篇梳理，每份材料单独成章。每章回答材料讲了什么、感知子系统如何设计、技术创新究竟在哪里、哪些结论可转化为工程动作。')
labeled(doc,'聚焦范围：','感知包含传感器布置与标定、锥桶/障碍物检测、相机-LiDAR 融合、时间同步、定位、SLAM、地图和健康监测；规划、控制论文仅说明其对感知输入的依赖。')
labeled(doc,'创新性判断：','算法论文说明方法的新意、所解问题和可量化结果；系统论文说明架构、冗余、同步和验证贡献；综述、规则与规划材料明确标为间接贡献，不把系统集成误写为通用算法创新。')
labeled(doc,'材料说明：','“00简介”和“2024FSAC入门指南”为参考材料；里斯本两份同题材料存在版本关系；TUM IAC 附件正文文本层不可读，已在对应章节如实标注。')
doc.add_heading('材料索引',2)
for i,(source,title,*_) in enumerate(items,1):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.15; font(p.add_run(f'{title}（源文件：{source}）'),9.5,False,(45,45,45))
for i,((source,title,overview,perception,takeaway),innovation) in enumerate(zip(items, innovations),1):
    doc.add_page_break(); doc.add_heading(f'第 {i:02d} 章  {title}',1)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(10); font(p.add_run('源文件：'+source),9.5,False,(90,90,90),True)
    doc.add_heading('内容概述',2); labeled(doc,'',overview)
    doc.add_heading('感知子系统要点',2); labeled(doc,'',perception)
    doc.add_heading('技术创新点与创新性判断',2); labeled(doc,'',innovation)
    doc.add_heading('工程启示与阅读结论',2); labeled(doc,'',takeaway)
doc.core_properties.title='无人方程式赛车论文感知子系统综述'; doc.core_properties.author='Codex'; doc.save(OUT)
print(OUT.resolve())
